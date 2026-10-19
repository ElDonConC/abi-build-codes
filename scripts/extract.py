#!/usr/bin/env python3
"""Extrae Catalogo_ABI_CORREGIDO_compatible.xlsx a JSON limpio en data/."""
import json
import pathlib
import sys
import unicodedata

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "Catalogo_ABI_CORREGIDO_compatible.xlsx"
OUT = ROOT / "data"

HEADER_ROW = 4  # encabezados en la fila 4, datos desde la 5

EXPECTED = {
    "Todas las builds": [
        "arma", "clase_de_arma", "nivel_economico", "precio_referencia_koen",
        "precio_minimo_publicado_koen", "precio_maximo_publicado_koen",
        "configuracion_etiqueta", "codigo_listo_para_copiar_prefijo_normalizado",
        "codigo_original_como_publicado", "estado_verificacion", "fuente_principal",
        "fecha_temporada", "url_fuente_principal", "otras_fuentes_donde_aparece",
        "notas_advertencias",
    ],
    "Armas maestras": [
        "arma", "clase", "estado_en_fuentes_de_referencia", "n_o_de_builds_codigos_unicos",
        "n_o_con_precio", "hay_barata", "hay_media", "hay_premium", "codigo_con_menor_precio",
        "precio_minimo_koen", "cobertura_economica", "fuente_del_roster", "url_de_referencia",
        "notas",
    ],
    "Armas sin build": [
        "arma", "clase", "estado_en_fuentes_de_referencia", "resultado", "fuente_del_roster",
        "url_de_referencia", "notas",
    ],
    "Lista aportada": [
        "arma_normalizada", "clase", "bloque_de_origen", "etiqueta_variante",
        "precio_en_lista_koen", "nivel_economico", "codigo_original_recibido",
        "codigo_listo_para_copiar", "cruce_exacto_con_catalogo_previo", "fuente_coincidente",
        "estado_de_verificacion_publica", "notas",
    ],
    "Fuentes y formato": [
        "fuente_publicacion", "tipo", "fecha_version", "que_aporta",
        "limitaciones_verificacion", "url_directa",
    ],
}


def slug(text):
    text = unicodedata.normalize("NFKD", str(text))
    text = "".join(c for c in text if not unicodedata.combining(c))
    keep = "".join(c if c.isalnum() else "_" for c in text.lower())
    while "__" in keep:
        keep = keep.replace("__", "_")
    return keep.strip("_")


def num(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    digits = "".join(c for c in str(value) if c.isdigit())
    return int(digits) if digits else None


def txt(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def main():
    if not XLSX.exists():
        sys.exit(f"No se encuentra {XLSX}")
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    def sheet_rows(name):
        ws = wb[name]
        ncols = len(EXPECTED[name])
        header = next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, max_col=ncols, values_only=True))
        headers = [slug(c) for c in header]
        missing = [c for c in EXPECTED[name] if c not in headers]
        if missing:
            sys.exit(f"Hoja '{name}': faltan columnas {missing}. Encabezados reales: {headers}")
        out = []
        for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            if not row or row[0] in (None, ""):
                continue
            out.append({k: (v if v != "" else None) for k, v in zip(headers, row)})
        return out

    builds = []
    for i, r in enumerate(sheet_rows("Todas las builds"), start=1):
        builds.append({
            "id": i,
            "arma": txt(r.get("arma")),
            "clase": txt(r.get("clase_de_arma")),
            "nivel": txt(r.get("nivel_economico")),
            "precio_ref": num(r.get("precio_referencia_koen")),
            "precio_min": num(r.get("precio_minimo_publicado_koen")),
            "precio_max": num(r.get("precio_maximo_publicado_koen")),
            "etiqueta": txt(r.get("configuracion_etiqueta")),
            "codigo": txt(r.get("codigo_listo_para_copiar_prefijo_normalizado")),
            "codigo_original": txt(r.get("codigo_original_como_publicado")),
            "estado": txt(r.get("estado_verificacion")),
            "fuente": txt(r.get("fuente_principal")),
            "fecha": txt(r.get("fecha_temporada")),
            "url_fuente": txt(r.get("url_fuente_principal")),
            "otras_fuentes": txt(r.get("otras_fuentes_donde_aparece")),
            "notas": txt(r.get("notas_advertencias")),
        })

    armas = []
    for i, r in enumerate(sheet_rows("Armas maestras"), start=1):
        armas.append({
            "id": i,
            "arma": txt(r.get("arma")),
            "clase": txt(r.get("clase")),
            "estado": txt(r.get("estado_en_fuentes_de_referencia")),
            "n_builds": num(r.get("n_o_de_builds_codigos_unicos")),
            "n_con_precio": num(r.get("n_o_con_precio")),
            "hay_barata": txt(r.get("hay_barata")),
            "hay_media": txt(r.get("hay_media")),
            "hay_premium": txt(r.get("hay_premium")),
            "codigo_mas_barato": txt(r.get("codigo_con_menor_precio")),
            "precio_min": num(r.get("precio_minimo_koen")),
            "cobertura": txt(r.get("cobertura_economica")),
            "fuente_roster": txt(r.get("fuente_del_roster")),
            "url_referencia": txt(r.get("url_de_referencia")),
            "notas": txt(r.get("notas")),
        })

    sin_build = []
    for i, r in enumerate(sheet_rows("Armas sin build"), start=1):
        sin_build.append({
            "id": i,
            "arma": txt(r.get("arma")),
            "clase": txt(r.get("clase")),
            "estado": txt(r.get("estado_en_fuentes_de_referencia")),
            "resultado": txt(r.get("resultado")),
            "fuente_roster": txt(r.get("fuente_del_roster")),
            "url_referencia": txt(r.get("url_de_referencia")),
            "notas": txt(r.get("notas")),
        })

    lista = []
    for i, r in enumerate(sheet_rows("Lista aportada"), start=1):
        lista.append({
            "id": i,
            "arma": txt(r.get("arma_normalizada")),
            "clase": txt(r.get("clase")),
            "bloque": txt(r.get("bloque_de_origen")),
            "etiqueta": txt(r.get("etiqueta_variante")),
            "precio": num(r.get("precio_en_lista_koen")),
            "nivel": txt(r.get("nivel_economico")),
            "codigo_original": txt(r.get("codigo_original_recibido")),
            "codigo": txt(r.get("codigo_listo_para_copiar")),
            "cruce": txt(r.get("cruce_exacto_con_catalogo_previo")),
            "fuente": txt(r.get("fuente_coincidente")),
            "estado": txt(r.get("estado_de_verificacion_publica")),
            "notas": txt(r.get("notas")),
        })

    fuentes = []
    for i, r in enumerate(sheet_rows("Fuentes y formato"), start=1):
        fuentes.append({
            "id": i,
            "nombre": txt(r.get("fuente_publicacion")),
            "tipo": txt(r.get("tipo")),
            "fecha": txt(r.get("fecha_version")),
            "aporta": txt(r.get("que_aporta")),
            "limitaciones": txt(r.get("limitaciones_verificacion")),
            "url": txt(r.get("url_directa")),
        })

    leeme = []
    for row in wb["LEEME"].iter_rows(min_row=5, values_only=True):
        if row and row[0]:
            leeme.append({"tema": txt(row[0]), "detalle": txt(row[1])})

    payload = {
        "builds": builds,
        "armas": armas,
        "sin_build": sin_build,
        "lista": lista,
        "fuentes": fuentes,
        "leeme": leeme,
    }

    OUT.mkdir(exist_ok=True)
    for name, data in payload.items():
        path = OUT / f"{name}.json"
        path.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{path.name}: {len(data)} registros")
    print("OK")


if __name__ == "__main__":
    main()
